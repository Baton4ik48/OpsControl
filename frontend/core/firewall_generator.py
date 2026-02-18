from openpyxl import load_workbook


def generate_firewall_rules(file_path: str) -> list[str]:
    wb = load_workbook(file_path)
    ws = wb.active

    if ws["A1"].value is None:
        raise ValueError("В A1 должен быть номер заявки")

    rule_name = str(ws["A1"].value).strip()

    commands = []
    rule_id = 1

    for row in range(3, ws.max_row + 1):
        src = ws[f"B{row}"].value
        dst = ws[f"F{row}"].value
        ports_raw = ws[f"E{row}"].value

        if not src or not dst or not ports_raw:
            continue

        ports = str(ports_raw).replace("\r", "").split("\n")

        port_parts = []
        for p in ports:
            p = p.strip()
            if ":" not in p:
                continue
            proto, port = p.split(":")
            port_parts.append(f"{proto.strip()} dport {port.strip()}")

        cmd = (
            f'firewall forward add {rule_id} '
            f'rule "{rule_name}" '
            f'src {src} dst {dst} '
            f'{" ".join(port_parts)} pass'
        )

        commands.append(cmd)
        rule_id += 1

    return commands
